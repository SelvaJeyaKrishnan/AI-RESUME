"""
Extracts raw text from uploaded resume files (PDF, DOCX, TXT).
Designed to fail gracefully: a broken/unsupported file raises ExtractionError
rather than crashing the request, so one bad resume never takes down a batch.
"""
import io

import fitz  # PyMuPDF
import docx  # python-docx


class ExtractionError(Exception):
    pass


def extract_text(contents: bytes, ext: str) -> str:
    ext = ext.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(contents)
        if ext == ".docx":
            return _extract_docx(contents)
        if ext == ".txt":
            return _extract_txt(contents)
        raise ExtractionError(f"Unsupported extension: {ext}")
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 - we want to catch anything a malformed file throws
        raise ExtractionError(f"Failed to parse {ext} file: {exc}") from exc


def _extract_pdf(contents: bytes) -> str:
    text_parts = []
    with fitz.open(stream=contents, filetype="pdf") as doc:
        if doc.is_encrypted:
            raise ExtractionError("PDF is password-protected")
        for page in doc:
            text_parts.append(page.get_text("text"))
    text = "\n".join(text_parts).strip()
    if not text:
        raise ExtractionError("No extractable text found (likely a scanned/image-only PDF)")
    return text


def _extract_docx(contents: bytes) -> str:
    document = docx.Document(io.BytesIO(contents))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(p for p in parts if p is not None).strip()
    if not text:
        raise ExtractionError("Document contains no readable text")
    return text


def _extract_txt(contents: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            text = contents.decode(encoding).strip()
            if text:
                return text
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode text file")
